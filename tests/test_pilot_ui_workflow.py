from __future__ import annotations

from datetime import UTC, date, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pymupdf
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pytest import MonkeyPatch

from ma_due_diligence.api.pilot_ai import (
    AiProviderPolicy,
    ControlledAiDraftService,
    LocalGroundedDraftProvider,
)
from ma_due_diligence.api.pilot_app import PilotApiContext, create_pilot_app
from ma_due_diligence.api.pilot_runs import PilotRunService
from ma_due_diligence.api.pilot_runtime import PilotApplicationRuntime
from ma_due_diligence.persistence import PilotDatabase, PilotDealRecord
from ma_due_diligence.security import DealRole
from ma_due_diligence.security.service import PilotAuthService
from ma_due_diligence.workflow_final import WorkflowApplication, build_workflow

_NOW = datetime(2026, 10, 25, 12, 0, tzinfo=UTC)


def _setup(tmp_path: Path, role: DealRole) -> tuple[TestClient, str, PilotDatabase]:
    database = PilotDatabase(tmp_path / "pilot.sqlite3")
    database.initialize()
    deal = PilotDealRecord(
        deal_id="deal-a",
        name="Project Northstar",
        target_name="SecureTech",
        buyer_name="Northstar Holdings",
        transaction_type="acquisition",
        reporting_currency="GBP",
        as_of_date=date(2026, 6, 30),
        workstreams=("financial", "commercial", "legal_contractual", "operational"),
        created_by="user-admin",
        created_at=_NOW,
        updated_at=_NOW,
    )
    database.create_deal(deal)
    auth = PilotAuthService(database)
    admin = auth.provision_user(
        email="admin@example.com",
        display_name="Pilot Admin",
        password="admin password long enough",
        is_pilot_admin=True,
        user_id="user-admin",
        occurred_at=_NOW,
    )
    member = auth.provision_user(
        email="member@example.com",
        display_name="Deal Team Member",
        password="member password long enough",
        user_id="user-member",
        occurred_at=_NOW,
    )
    principal = auth.login(
        email=admin.email,
        password="admin password long enough",
    ).principal
    auth.assign_deal_role(
        actor=principal,
        deal_id=deal.deal_id,
        user_id=member.user_id,
        role=role,
    )
    runs = PilotRunService(database, WorkflowApplication(build_workflow()))
    ai = ControlledAiDraftService(
        database,
        LocalGroundedDraftProvider(),
        AiProviderPolicy(
            provider_id="local-controlled-demo",
            model_id="grounded-template-v1",
            external_processing=False,
            retention_policy="test deal retention policy",
            training_use=False,
        ),
    )
    client = TestClient(create_pilot_app(PilotApiContext(database, auth, runs, ai=ai)))
    login = client.post(
        "/pilot/auth/login",
        json={"email": member.email, "password": "member password long enough"},
    )
    return client, str(login.json()["access_token"]), database


def _headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def test_packaged_ui_is_served_with_accessible_sign_in(tmp_path: Path) -> None:
    client, _, _ = _setup(tmp_path, DealRole.REVIEWER)

    response = client.get("/")
    script = client.get("/assets/app.js")
    styles = client.get("/assets/styles.css")

    assert response.status_code == 200
    assert "DealLens" in response.text
    assert "Sign in to your deal workspace" in response.text
    assert script.status_code == 200
    assert styles.status_code == 200


def test_runtime_bootstraps_configured_administrator(
    tmp_path: Path,
    monkeypatch: MonkeyPatch,
) -> None:
    monkeypatch.setenv(
        "MADD_CHECKPOINT_KEY",
        "AAECAwQFBgcICQoLDA0ODxAREhMUFRYXGBkaGxwdHh8=",
    )
    monkeypatch.setenv(
        "MADD_DOCUMENT_KEY",
        "AAECAwQFBgcICQoLDA0ODxAREhMUFRYXGBkaGxwdHh8=",
    )
    monkeypatch.setenv("MADD_BOOTSTRAP_ADMIN_EMAIL", "pilot@example.com")
    monkeypatch.setenv("MADD_BOOTSTRAP_ADMIN_PASSWORD", "pilot password long enough")
    monkeypatch.setenv("MADD_BOOTSTRAP_ADMIN_NAME", "Pilot Reviewer")
    runtime = PilotApplicationRuntime(tmp_path / "pilot-data")
    try:
        response = TestClient(runtime.app).post(
            "/pilot/auth/login",
            json={
                "email": "pilot@example.com",
                "password": "pilot password long enough",
            },
        )
    finally:
        runtime.close()

    assert response.status_code == 200
    assert response.json()["user"]["is_pilot_admin"] is True


def test_reviewer_can_run_review_and_retrieve_report_through_pilot_api(tmp_path: Path) -> None:
    client, token, _ = _setup(tmp_path, DealRole.REVIEWER)
    headers = _headers(token)

    started = client.post("/pilot/deals/deal-a/runs", headers=headers)
    assert started.status_code == 201
    paused = started.json()
    assert paused["status"] == "review_required"
    assert paused["document_count"] == 25
    assert paused["finding_count"] == 14
    assert paused["findings"]
    assert paused["financial"]["adjusted_ebitda"] == "15"

    decisions = [
        {
            "subject_type": "finding",
            "subject_id": finding_id,
            "action": "approve_finding",
            "rationale": "Reviewed supporting evidence.",
        }
        for finding_id in [item["finding_id"] for item in paused["findings"]]
    ]
    for conflict in paused["conflicts"]:
        if conflict["conflict_id"] in paused["review_conflict_ids"]:
            decisions.append(
                {
                    "subject_type": "conflict",
                    "subject_id": conflict["conflict_id"],
                    "action": "resolve_conflict",
                    "rationale": "Selected direct source for presentation.",
                    "preferred_source_id": conflict["observation_fact_ids"][0],
                }
            )
    reviewed = client.post(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/review",
        headers=headers,
        json={"rationale": "Completed material review.", "decisions": decisions},
    )
    assert reviewed.status_code == 200
    assert reviewed.json()["report_available"] is True
    assert reviewed.json()["review_action_count"] == len(decisions)

    history = client.get(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/review-history",
        headers=headers,
    )
    assert history.status_code == 200
    events = history.json()["events"]
    assert history.json()["event_count"] == 1 + (2 * len(decisions))
    assert events[0]["event_type"] == "commented"
    assert {item["event_type"] for item in events} >= {
        "assigned",
        "approved",
        "conflict_resolved",
    }
    decided_finding_ids = {
        item["subject_id"]
        for item in events
        if item["subject_type"] == "finding" and item["event_type"] == "approved"
    }
    assert decided_finding_ids == {item["finding_id"] for item in paused["findings"]}

    report = client.get(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/report",
        headers=headers,
    )
    assert report.status_code == 200
    assert report.json()["report"]["target_name"] == "SecureTech"


def test_review_rejects_partial_finding_disposition_without_history(tmp_path: Path) -> None:
    client, token, _ = _setup(tmp_path, DealRole.REVIEWER)
    headers = _headers(token)
    paused = client.post("/pilot/deals/deal-a/runs", headers=headers).json()

    response = client.post(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/review",
        headers=headers,
        json={
            "rationale": "Incomplete review should not release a report.",
            "decisions": [
                {
                    "subject_type": "finding",
                    "subject_id": paused["findings"][0]["finding_id"],
                    "action": "approve_finding",
                    "rationale": "Only one finding was reviewed.",
                }
            ],
        },
    )

    assert response.status_code == 400
    assert "every released finding" in response.json()["detail"]
    history = client.get(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/review-history",
        headers=headers,
    )
    assert history.json() == {"events": [], "event_count": 0}


def test_edit_reject_severity_and_information_request_events_are_retained(
    tmp_path: Path,
) -> None:
    client, token, _ = _setup(tmp_path, DealRole.REVIEWER)
    headers = _headers(token)
    paused = client.post("/pilot/deals/deal-a/runs", headers=headers).json()
    actions = [
        ("reject_finding", {}),
        ("change_severity", {"severity": "medium"}),
        ("mark_nonmaterial", {}),
        ("request_more_evidence", {}),
    ]
    decisions = []
    for index, finding in enumerate(paused["findings"]):
        action, extra = actions[index] if index < len(actions) else ("approve_finding", {})
        decisions.append(
            {
                "subject_type": "finding",
                "subject_id": finding["finding_id"],
                "action": action,
                "rationale": f"Reviewer disposition for finding {index + 1}.",
                **extra,
            }
        )
    for conflict in paused["conflicts"]:
        if conflict["conflict_id"] in paused["review_conflict_ids"]:
            decisions.append(
                {
                    "subject_type": "conflict",
                    "subject_id": conflict["conflict_id"],
                    "action": "resolve_conflict",
                    "rationale": "Selected the controlling source.",
                    "preferred_source_id": conflict["observation_fact_ids"][0],
                }
            )

    reviewed = client.post(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/review",
        headers=headers,
        json={"rationale": "Completed independent dispositions.", "decisions": decisions},
    )

    assert reviewed.status_code == 200
    assert reviewed.json()["status"] == "waiting_for_information"
    assert reviewed.json()["report_available"] is False
    assert reviewed.json()["findings"][0]["status"] == "closed"
    assert reviewed.json()["findings"][1]["severity"] == "medium"
    history = client.get(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/review-history",
        headers=headers,
    ).json()
    event_types = {item["event_type"] for item in history["events"]}
    assert event_types >= {
        "assigned",
        "commented",
        "rejected",
        "severity_changed",
        "edited",
        "information_requested",
        "conflict_resolved",
    }


def test_analyst_can_run_analysis_but_cannot_submit_review(tmp_path: Path) -> None:
    client, token, _ = _setup(tmp_path, DealRole.ANALYST)
    headers = _headers(token)
    started = client.post("/pilot/deals/deal-a/runs", headers=headers)
    run_id = str(started.json()["run_id"])

    review = client.post(
        f"/pilot/deals/deal-a/runs/{run_id}/review",
        headers=headers,
        json={
            "rationale": "Attempted review.",
            "decisions": [
                {
                    "subject_type": "finding",
                    "subject_id": started.json()["review_finding_ids"][0],
                    "action": "approve_finding",
                    "rationale": "Not authorized.",
                }
            ],
        },
    )

    assert review.status_code == 403


def test_controlled_ai_draft_requires_explicit_human_use_and_retains_provenance(
    tmp_path: Path,
) -> None:
    client, token, database = _setup(tmp_path, DealRole.REVIEWER)
    headers = _headers(token)
    paused = client.post("/pilot/deals/deal-a/runs", headers=headers).json()
    first_finding = paused["findings"][0]

    policy = client.get("/pilot/deals/deal-a/ai-policy", headers=headers)
    draft = client.post(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/ai-drafts/findings/{first_finding['finding_id']}",
        headers=headers,
    )

    assert policy.status_code == 200
    assert policy.json()["human_approval_required"] is True
    assert policy.json()["may_change_financials"] is False
    assert draft.status_code == 201
    draft_payload = draft.json()
    assert draft_payload["approval_required"] is True
    assert draft_payload["applied"] is False
    assert set(draft_payload["citation_ids"]) <= {
        item["evidence_id"] for item in first_finding["evidence"]
    }
    unchanged = client.get(f"/pilot/deals/deal-a/runs/{paused['run_id']}", headers=headers).json()
    assert unchanged["status"] == "review_required"
    assert unchanged["review_action_count"] == 0

    decisions = []
    for index, finding in enumerate(paused["findings"]):
        decision = {
            "subject_type": "finding",
            "subject_id": finding["finding_id"],
            "action": "approve_finding",
            "rationale": "Reviewed cited evidence independently.",
        }
        if index == 0:
            decision["rationale"] = draft_payload["narrative"]
            decision["ai_draft_id"] = draft_payload["draft_id"]
        decisions.append(decision)
    decisions.extend(
        {
            "subject_type": "conflict",
            "subject_id": conflict["conflict_id"],
            "action": "resolve_conflict",
            "rationale": "Selected the controlling source independently.",
            "preferred_source_id": conflict["observation_fact_ids"][0],
        }
        for conflict in paused["conflicts"]
        if conflict["conflict_id"] in paused["review_conflict_ids"]
    )
    reviewed = client.post(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/review",
        headers=headers,
        json={
            "rationale": "Completed human review after inspecting the optional draft.",
            "decisions": decisions,
        },
    )

    assert reviewed.status_code == 200
    history = client.get(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/review-history",
        headers=headers,
    ).json()["events"]
    assisted = next(
        item
        for item in history
        if item["subject_id"] == first_finding["finding_id"] and item["event_type"] == "approved"
    )
    assert assisted["metadata"]["ai_assisted"] is True
    assert assisted["metadata"]["ai_draft_id"] == draft_payload["draft_id"]
    artifacts = database.list_artifacts("deal-a", paused["run_id"])
    assert artifacts[0].payload["approval_status"] == "pending_human_review"
    audit_actions = [item.action for item in database.list_audit_events("deal-a")]
    assert "ai_draft_generated" in audit_actions
    review_audit = next(
        item for item in database.list_audit_events("deal-a") if item.action == "review_submitted"
    )
    assert review_audit.detail["ai_assisted_decision_count"] == 1


def test_ai_draft_from_another_finding_cannot_be_attached(tmp_path: Path) -> None:
    client, token, _ = _setup(tmp_path, DealRole.REVIEWER)
    headers = _headers(token)
    paused = client.post("/pilot/deals/deal-a/runs", headers=headers).json()
    draft = client.post(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/ai-drafts/findings/{paused['findings'][0]['finding_id']}",
        headers=headers,
    ).json()
    decisions = [
        {
            "subject_type": "finding",
            "subject_id": finding["finding_id"],
            "action": "approve_finding",
            "rationale": "Reviewed independently.",
            **({"ai_draft_id": draft["draft_id"]} if index == 1 else {}),
        }
        for index, finding in enumerate(paused["findings"])
    ]
    decisions.extend(
        {
            "subject_type": "conflict",
            "subject_id": conflict["conflict_id"],
            "action": "resolve_conflict",
            "rationale": "Selected controlling source.",
            "preferred_source_id": conflict["observation_fact_ids"][0],
        }
        for conflict in paused["conflicts"]
        if conflict["conflict_id"] in paused["review_conflict_ids"]
    )

    response = client.post(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/review",
        headers=headers,
        json={"rationale": "Attempted mismatched draft use.", "decisions": decisions},
    )

    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "policy_blocked"
    unchanged = client.get(f"/pilot/deals/deal-a/runs/{paused['run_id']}", headers=headers).json()
    assert unchanged["status"] == "review_required"


def test_released_report_and_workbook_exports_match_application_state(tmp_path: Path) -> None:
    client, token, database = _setup(tmp_path, DealRole.REVIEWER)
    headers = _headers(token)
    paused = client.post("/pilot/deals/deal-a/runs", headers=headers).json()
    decisions = [
        {
            "subject_type": "finding",
            "subject_id": finding["finding_id"],
            "action": "approve_finding",
            "rationale": f"Approved {finding['finding_id']} against cited evidence.",
        }
        for finding in paused["findings"]
    ]
    decisions.extend(
        {
            "subject_type": "conflict",
            "subject_id": conflict["conflict_id"],
            "action": "resolve_conflict",
            "rationale": "Selected the controlling direct source.",
            "preferred_source_id": conflict["observation_fact_ids"][0],
        }
        for conflict in paused["conflicts"]
        if conflict["conflict_id"] in paused["review_conflict_ids"]
    )
    reviewed = client.post(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/review",
        headers=headers,
        json={"rationale": "Completed export release review.", "decisions": decisions},
    ).json()
    report_payload = client.get(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/report",
        headers=headers,
    ).json()["report"]

    pdf = client.get(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/exports/report.pdf",
        headers=headers,
    )
    workbook_response = client.get(
        f"/pilot/deals/deal-a/runs/{paused['run_id']}/exports/workbook.xlsx",
        headers=headers,
    )

    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.headers["cache-control"] == "no-store"
    assert pdf.headers["x-content-type-options"] == "nosniff"
    assert pdf.headers["x-export-sha256"] == sha256(pdf.content).hexdigest()
    assert "SecureTech-reviewed-diligence-report.pdf" in pdf.headers["content-disposition"]
    with pymupdf.open(stream=pdf.content, filetype="pdf") as document:  # type: ignore[no-untyped-call]
        pdf_text = "\n".join(page.get_text() for page in document)
    assert "Reviewed M&A Due-Diligence Report" in pdf_text
    assert "Diligence-adjusted EBITDA" in pdf_text
    assert "GBP 15m" in pdf_text
    assert "Explicit limitations" in pdf_text
    assert report_payload["limitations"][0] in pdf_text

    assert workbook_response.status_code == 200
    assert workbook_response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        workbook_response.headers["x-export-sha256"]
        == sha256(workbook_response.content).hexdigest()
    )
    workbook = load_workbook(BytesIO(workbook_response.content), data_only=False)
    assert workbook.sheetnames == [
        "Read Me",
        "Findings Register",
        "Citations",
        "Financial Summary",
        "EBITDA Bridge",
        "Working Capital",
        "NWC Peg",
        "Customer Concentration",
        "Net Debt",
        "Revenue Reconciliation",
        "Review History",
        "Limitations",
    ]
    findings_sheet = workbook["Findings Register"]
    assert findings_sheet.max_row == reviewed["finding_count"] + 1
    finding_headers = [cell.value for cell in findings_sheet[1]]
    first_finding = dict(
        zip(finding_headers, [cell.value for cell in findings_sheet[2]], strict=True)
    )
    assert first_finding["Finding ID"] == reviewed["findings"][0]["finding_id"]
    assert first_finding["Severity"] == reviewed["findings"][0]["severity"]
    assert first_finding["Review status"] == reviewed["findings"][0]["analyst_review_status"]
    assert first_finding["Decision"] == "approved"
    summary = {
        row[0].value: row[1].value for row in workbook["Financial Summary"].iter_rows(min_row=2)
    }
    assert summary["Diligence-adjusted EBITDA"] == 15
    assert summary["Normalized NWC"] == float(reviewed["financial"]["normalized_nwc"])
    limitations = [row[1].value for row in workbook["Limitations"].iter_rows(min_row=2)]
    assert limitations == report_payload["limitations"]
    assert workbook["Review History"].max_row == 2 + (2 * len(decisions))
    audit = database.list_audit_events("deal-a")
    assert [event.action for event in audit[-2:]] == ["report_exported", "workbook_exported"]
    assert audit[-1].detail["sha256"] == workbook_response.headers["x-export-sha256"]
    assert "content" not in audit[-1].detail


def test_exports_fail_closed_before_release_and_across_deals(tmp_path: Path) -> None:
    client, token, _ = _setup(tmp_path, DealRole.REVIEWER)
    headers = _headers(token)
    paused = client.post("/pilot/deals/deal-a/runs", headers=headers).json()
    path = f"/pilot/deals/deal-a/runs/{paused['run_id']}/exports/report.pdf"

    assert client.get(path, headers=headers).status_code == 409
    assert client.get(path).status_code == 401
    assert (
        client.get(
            f"/pilot/deals/deal-b/runs/{paused['run_id']}/exports/report.pdf",
            headers=headers,
        ).status_code
        == 404
    )
